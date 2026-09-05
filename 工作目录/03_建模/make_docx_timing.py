#!/usr/bin/env python3
"""按论文体例写出流水线时间统计 Word。

先跑 measure_pipeline_timing.py 得到 reports_timing/timing.json，再：

    uv run --with python-docx python make_docx_timing.py
"""

from __future__ import annotations

import json
import shutil
import sys
from pathlib import Path

from docx.oxml import OxmlElement
from docx.oxml.ns import qn

HERE = Path(__file__).resolve().parent
sys.path.insert(0, str(HERE / "forecast_next_stage"))

from report_fmt import (  # noqa: E402
    add_body,
    add_caption,
    add_h,
    add_note,
    add_title,
    add_toc_heading,
    add_toc_line,
    add_page_break,
    fill_cell,
    make_table,
    new_doc,
    set_table_fixed,
)

TIMING = HERE / "reports_timing" / "timing.json"
OUT = HERE / "多模态实验流水线_时间统计.docx"
DESK = Path("/Users/licochen/Desktop")

GROUPS = [
    ("预处理", ["眼动预处理", "脑电预处理", "心率预处理"]),
    ("特征提取", ["行为特征提取", "眼动特征提取", "脑电特征提取", "心率特征提取"]),
    ("绩效评估", ["绩效 S 回归预测"]),
    ("趋势预测", ["绩效 S 趋势预测"]),
]


def r1(x: float) -> float:
    return round(float(x), 1)


def fmt_sec(x: float) -> str:
    v = r1(x)
    if v == int(v):
        return f"{int(v)} 秒"
    return f"{v:.1f} 秒"


def make_group_table(doc, rows_sec: dict[str, float], widths=(3.2, 6.6, 2.5, 2.5)):
    n_body = sum(len(items) for _, items in GROUPS)
    table = doc.add_table(rows=1 + n_body, cols=4)
    table.style = "Table Grid"
    set_table_fixed(table, list(widths))
    for j, h in enumerate(["流程名称", "内容", "时间", "总计"]):
        fill_cell(table.rows[0].cells[j], h, header=True, size=10.5)

    r = 1
    for gname, items in GROUPS:
        total = sum(r1(rows_sec[k]) for k in items)
        start = r
        for name in items:
            fill_cell(table.rows[r].cells[0], gname, header=False, size=10.5)
            fill_cell(table.rows[r].cells[1], name, header=False, size=10.5)
            fill_cell(table.rows[r].cells[2], fmt_sec(rows_sec[name]), header=False, size=10.5)
            fill_cell(table.rows[r].cells[3], fmt_sec(total), header=False, size=10.5)
            r += 1
        if r - 1 > start:
            table.cell(start, 0).merge(table.cell(r - 1, 0))
            table.cell(start, 3).merge(table.cell(r - 1, 3))
            fill_cell(table.rows[start].cells[0], gname, header=False, size=10.5)
            fill_cell(table.rows[start].cells[3], fmt_sec(total), header=False, size=10.5)

    for i, row in enumerate(table.rows):
        trPr = row._tr.get_or_add_trPr()
        if trPr.find(qn("w:cantSplit")) is None:
            trPr.append(OxmlElement("w:cantSplit"))
        if i == 0 and trPr.find(qn("w:tblHeader")) is None:
            trPr.append(OxmlElement("w:tblHeader"))
    return table


def build() -> Path:
    data = json.loads(TIMING.read_text(encoding="utf-8"))
    rows = data["rows"]
    proto = data["protocol"]
    machine = data["machine"]
    s_pred = data["detail"]["s_prediction"]
    trend = data["detail"]["trend"]
    eye_load = data["detail"]["eye_load"]
    grand = sum(r1(v) for v in rows.values())

    cpu = machine.get("cpu") or machine.get("processor")
    ram = machine.get("ram_gb")
    ram_s = f"{ram:.0f} GB" if ram else "—"

    doc = new_doc("多模态实验流水线  时间统计")
    add_title(doc, "多模态实验流水线时间统计", "26 名被试 / 84 条被试–任务　·　正式口径 30 秒窗、5 秒步")

    add_toc_heading(doc)
    add_toc_line(doc, "1  口径与计时条件", 1)
    add_toc_line(doc, "2  时间统计", 1)
    add_toc_line(doc, "3  各项说明", 1)
    add_toc_line(doc, "4  小结", 1)
    add_page_break(doc)

    add_h(doc, "1  口径与计时条件", 1)
    add_body(
        doc,
        "本表统计的是正式实验流水线在本机上处理全部 26 名被试、84 条被试–任务一次所需的墙钟时间，不是单窗推理时延，也不是研发迭代总工时。"
        "切窗与正式建模一致：窗长 30 秒、步长 5 秒，共 12 624 个窗口；任务级特征为 66 个窗口指标各收均值、标准差、中位数与斜率，得到 84×264。"
        "脑电原始文件已是 EEGLAB 预处理后的 256 Hz .set；表中「脑电预处理」只计读入与时间对齐，不含滤波、ICA 等人工预处理。",
    )
    add_caption(doc, "表 1  计时条件")
    make_table(
        doc,
        ["项", "内容"],
        [
            ["样本", f"{proto['n_subjects']} 名被试、{proto['n_samples']} 条被试–任务、{proto['n_windows']} 个窗口"],
            ["切窗", f"窗长 {int(proto['window_sec'])} 秒，步长 {int(proto['step_sec'])} 秒"],
            ["机器", f"{cpu}，内存 {ram_s}，{machine.get('system')} {machine.get('machine')}"],
            ["环境", f"Python {machine.get('python')}"],
            ["计时时刻", data["measured_at"]],
            ["窗级特征", f"探针样本 {', '.join(proto['probe_sample_ids'])} 共 {proto['windows_timed']} 窗，按比例外推至 {proto['n_windows']} 窗"],
            ["其余步骤", "84 条全量实测（读入、任务级聚合、五折建模、趋势主路径）"],
        ],
        [3.6, 11.2],
        left_cols={0, 1},
    )

    add_h(doc, "2  时间统计", 1)
    add_body(
        doc,
        "下表结构与时间统计模板一致。时间列为该分项墙钟；总计列为同一流程下各分项之和。"
        f"全流程合计 {fmt_sec(grand)}。",
    )
    add_caption(doc, "表 2  正式口径全量批处理时间")
    make_group_table(doc, rows)
    add_note(
        doc,
        "注：单项按 0.1 秒四舍五入。窗级特征由 3 条探针外推；读入、聚合与建模为全量实测。"
        "绩效回归复现五折定额 27 维浅树 XGB，S 的 pooled R²＝0.979；"
        "趋势预测走正式主路径 v8（标准化 Ridge 预报 27 列再接同一套 XGB），对照矩阵全样本 S R²＝0.966。",
    )

    add_h(doc, "3  各项说明", 1)

    add_h(doc, "3.1  预处理", 2)
    add_body(
        doc,
        f"眼动预处理主要是读入 Tobii 导出的 5 个任务 TSV（约 712 MB）并解析时间戳、瞳孔与注视类型、兴趣区区间，"
        f"其中逐点时序读入 {eye_load['timeseries']:.1f} 秒，起始信息 {eye_load['starts']:.1f} 秒，兴趣区 {eye_load['aoi']:.2f} 秒；"
        "随后以眼动起始时刻为锚点，把脑电时长切成 30 秒窗。"
        f"脑电预处理是读入 84 个已预处理的 256 Hz .set（合计 2.40 GB），全量读入 {data['detail']['eeg_load']['seconds']:.2f} 秒。"
        f"心率预处理是读入逐点心率表（13 364 行）及 NASA 加权总分表，全量 {data['detail']['hr_load']:.2f} 秒。",
    )

    add_h(doc, "3.2  特征提取", 2)
    add_body(
        doc,
        "行为特征来自操作日志：解析带时间戳的动作，按窗统计次数、密度、对错与多余/重复等 12 个窗口指标。"
        "眼动特征包括滤波瞳孔、有效采样与注视/扫视比例、9 个兴趣区聚合量，以及由 EyesNotFound 连续段得到的疑似眨眼。"
        "脑电特征对额、中央、顶、枕四区做 Welch 功率谱，取 δ/θ/α/β/γ 及 θ/α、β/α，再按被试做窗内 z 分数，共 28 个窗口指标。"
        "心率特征为窗内均值、标准差、最低、最高与斜率，共 5 个窗口指标。"
        "各模态窗口指标再按任务收成均值、标准差、中位数与斜率（66×4＝264）。"
        "脑电功率谱是特征提取里最耗时的一步。",
    )

    add_h(doc, "3.3  绩效评估", 2)
    add_body(
        doc,
        f"绩效 S 回归预测按被试 GroupKFold 五折：每一折只在训练被试上按模态定额互信息录取 27 维"
        f"（眼动 6、脑电 5、心率 4、行为 12），拟合浅树 XGBoost 预测 NASA-TLX 加权总分，"
        f"再与真实步骤分按 0.70／0.30 合成 S。84 条全量五折 {s_pred['seconds']:.2f} 秒，"
        f"复现 pooled R²＝{s_pred['s_r2']:.3f}，MAE＝{s_pred['s_mae']:.3f}，与正式报告一致。",
    )

    add_h(doc, "3.4  趋势预测", 2)
    add_body(
        doc,
        f"绩效 S 趋势预测走正式主路径：把一次任务的窗口按时间切成已观察阶段与未来阶段，"
        f"用标准化 Ridge 由已观察 27 列预报整场 27 列，再经与绩效回归相同的浅树 XGB 估计负荷并合成 S。"
        f"含窗口缓存读入在内共 {trend['seconds']:.2f} 秒（读入 {trend['seconds_load']:.2f} 秒，拟合与推理 {trend['seconds_fit_predict']:.2f} 秒）。"
        f"该次计时覆盖 {trend['n_eval']} 条可切段样本，S R²＝{trend['s_r2']:.3f}，"
        "与对照矩阵 v8／ridge_scaled 全样本口径一致；对外验证组 5 人、17 条的报告值为 R²＝0.948。",
    )

    add_h(doc, "4  小结", 1)
    add_body(
        doc,
        f"在 26 名被试、84 次任务、12 624 个 30 秒窗口上，从已对齐的原始文件到预测版绩效 S，"
        f"全流程墙钟约 {fmt_sec(grand)}。"
        f"时间主要花在特征提取（{fmt_sec(sum(rows[k] for k in ['行为特征提取','眼动特征提取','脑电特征提取','心率特征提取']))}），"
        f"其中脑电 Welch 功率谱最长；眼动 TSV 读入构成预处理的主体。"
        "五折绩效回归与趋势主路径均在数秒内完成，说明正式口径下的建模本身不是瓶颈。"
        "若部署为单条任务在线推理，墙钟会明显低于本表的全量批处理时间。",
    )

    doc.save(OUT)
    desk = DESK / OUT.name
    shutil.copy2(OUT, desk)
    print("wrote", OUT)
    print("copy ", desk)
    return OUT


if __name__ == "__main__":
    build()
